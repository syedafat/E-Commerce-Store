from datetime import datetime, timedelta, timezone
import io
from pathlib import Path
import re

from bson import ObjectId
from openpyxl import load_workbook

from tests.conftest import signed_in


def test_forecast_uses_dated_records_and_reports_error():
    from forecasting import forecast_sales
    start = datetime(2025, 1, 6, tzinfo=timezone.utc)
    records = [{
        "sale_date": start + timedelta(days=7 * i), "total_revenue": 100 + i,
        "profit_per_sale": 25 + i,
    } for i in range(10)]
    records[0]["sale_date"] = "2025-01-06T00:00:00-05:00"
    records.append({"sale_date": "not-a-date", "total_revenue": 10, "profit_per_sale": 1})
    result = forecast_sales(records)
    assert result["available"] is True
    assert result["horizon"] == "next 4 weeks"
    assert result["mae"] is not None
    assert result["revenue"] > 0
    assert result["skipped_records"] == 1


def test_forecast_refuses_insufficient_data():
    from forecasting import forecast_sales
    assert forecast_sales([])["available"] is False


def test_forecast_rejects_history_without_recent_sales_weeks():
    from forecasting import forecast_sales
    start = datetime(2025, 1, 6, tzinfo=timezone.utc)
    records = [{"sale_date": start + timedelta(days=7 * i), "total_revenue": 10, "profit_per_sale": 2}
               for i in (0, 1, 9)]
    result = forecast_sales(records)
    assert result["available"] is False
    assert "four sales weeks" in result["message"]


def test_all_admin_pages_render(client):
    from routes import product_management
    product_id = ObjectId()
    product_management.product_collection.insert_one({
        "_id": product_id, "product_name": "Widget", "category": "Misc",
        "price": 10, "selling_price": 10, "cost_price": 4,
        "available_quantity": 2, "description": "Test item",
    })
    signed_in(client)
    for path in ("/admin", "/products", f"/product-info/{product_id}",
                 "/product-information", "/inventory", "/inventory-details",
                 "/sales", "/realtime_prediction", "/reports", "/add-product",
                 f"/edit-product/{product_id}"):
        response = client.get(path)
        assert response.status_code == 200, (path, response.status_code, response.get_data(as_text=True)[:500])


def test_forecast_route_link_is_live(client):
    signed_in(client)
    response = client.get("/realtime_prediction")
    assert response.status_code == 200
    assert b"Forecast unavailable" in response.data


def test_product_create_edit_delete_uses_canonical_collection(client):
    from routes import product_management
    signed_in(client)
    payload = {"csrf_token": "test-csrf", "name": "Lamp", "category": "Home",
               "price": "12.50", "cost_price": "5", "quantity": "4", "description": "Desk lamp"}
    assert client.post("/add-product", data=payload).status_code == 302
    product = product_management.product_collection.find_one({"product_name": "Lamp"})
    assert product["available_quantity"] == 4
    payload.update({"name": "Desk lamp", "price": "13", "quantity": "3"})
    assert client.post(f"/edit-product/{product['_id']}", data=payload).status_code == 302
    edited = product_management.product_collection.find_one({"_id": product["_id"]})
    assert edited["product_name"] == "Desk lamp"
    assert edited["available_quantity"] == 3
    assert client.post(f"/delete-product/{product['_id']}", data={"csrf_token": "test-csrf"}).status_code == 302
    archived = product_management.product_collection.find_one({"_id": product["_id"]})
    assert archived["is_archived"] is True


def test_ai_export_uses_database_derived_forecast(client):
    from routes import reports
    start = datetime(2025, 1, 6, tzinfo=timezone.utc)
    reports.sales_collection.insert_many([{
        "sale_date": start + timedelta(days=7 * i), "total_revenue": 100 + i,
        "profit_per_sale": 20 + i,
    } for i in range(10)])
    signed_in(client)
    response = client.get("/reports/export?type=ai&format=excel")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    rows = list(workbook.active.values)
    assert rows[0][0] == "available"
    assert rows[1][0] is True
    assert rows[1][2] > 0


def test_excel_export_escapes_formula_like_user_data(client):
    from routes import reports
    reports.product_info_collection.insert_one({
        "product_name": "=HYPERLINK(\"https://attacker.invalid\")",
        "category": "Home", "price": 5, "available_quantity": 1,
    })
    signed_in(client)
    response = client.get("/reports/export?type=inventory&format=excel")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True, data_only=False)
    headers = list(next(workbook.active.iter_rows(values_only=True)))
    row = list(next(workbook.active.iter_rows(min_row=2, values_only=True)))
    value = row[headers.index("product_name")]
    assert value.startswith("'=")


def test_pnl_report_totals_are_calculated_from_all_sales(client):
    from routes import reports
    reports.sales_collection.insert_many([
        {"total_revenue": 20, "discount_applied": 1, "profit_per_sale": 5},
        {"total_revenue": 30, "discount_applied": 2, "profit_per_sale": -3},
    ])
    signed_in(client)
    response = client.get("/reports/export?type=pnl&format=excel")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data), read_only=True)
    rows = list(workbook.active.values)
    assert rows[1] == (2, 50, 3, 2)


def test_pnl_report_refuses_malformed_financial_values(client):
    from routes import reports
    reports.sales_collection.insert_one({"total_revenue": "unknown", "profit_per_sale": 2})
    signed_in(client)
    response = client.get("/reports/export?type=pnl&format=excel")
    assert response.status_code == 400
    assert b"invalid values" in response.data


def test_product_search_filter_and_pagination(client):
    from routes import product_management
    signed_in(client)
    product_management.product_collection.insert_many([
        {"product_name": f"Lamp {i}", "category": "Home", "price": i+1, "available_quantity": i+1}
        for i in range(30)
    ])
    page = client.get("/products?category=Home&sort=price&page=2")
    assert page.status_code == 200
    assert b"Page 2 of 2" in page.data
    filtered = client.get("/products?q=Lamp%202&stock=available")
    assert filtered.status_code == 200
    assert b"Lamp 2" in filtered.data


def test_product_with_sales_history_cannot_be_deleted(client):
    from routes import product_management
    product_id = product_management.product_collection.insert_one({
        "product_name": "Historical", "category": "Home", "price": 4,
        "available_quantity": 0,
    }).inserted_id
    product_management.sales_collection.insert_one({"product_id": product_id, "product_name": "Historical"})
    signed_in(client)
    response = client.post(f"/delete-product/{product_id}", data={"csrf_token": "test-csrf"})
    assert response.status_code == 302
    assert product_management.product_collection.find_one({"_id": product_id}) is not None


def test_product_default_percentage_discount_is_reflected_in_catalog_price(client):
    from routes import product_management
    signed_in(client)
    response = client.post("/add-product", data={
        "csrf_token": "test-csrf", "product_name": "Marked down", "category": "Home",
        "selling_price": "100", "cost_price": "40", "discount_percentage": "10",
        "available_quantity": "2", "reorder_level": "1",
    })
    assert response.status_code == 302
    product = product_management.product_collection.find_one({"product_name": "Marked down"})
    assert product["selling_price"] == 100
    assert product["price"] == 90


def test_csv_exports_escape_formula_like_values_and_apply_sale_dates(client):
    from routes import reports
    start = datetime(2025, 3, 1, tzinfo=timezone.utc)
    reports.sales_collection.insert_many([
        {"product_name": "=cmd", "sale_date": start, "total_revenue": 12, "quantity_sold": 1},
        {"product_name": "Other", "sale_date": start-timedelta(days=1), "total_revenue": 8, "quantity_sold": 1},
    ])
    signed_in(client)
    response = client.get("/reports/export?type=sales&format=csv&start=2025-03-01&end=2025-03-01")
    assert response.status_code == 200
    assert b"'=cmd" in response.data
    assert b"Other" not in response.data


def test_inventory_adjustments_create_movement_history(client):
    from routes import inventory
    product_id = inventory.product_info_collection.insert_one({
        "product_name": "Stock item", "available_quantity": 4, "reorder_level": 2,
    }).inserted_id
    signed_in(client)
    response = client.post(f"/inventory/update-stock/{product_id}", data={
        "csrf_token": "test-csrf", "quantity": "3", "update_type": "add",
        "reorder_level": "6", "reason": "Supplier receipt",
    })
    assert response.status_code == 302
    product = inventory.product_info_collection.find_one({"_id": product_id})
    movement = inventory.inventory_movements_collection.find_one({"product_id": product_id})
    assert product["available_quantity"] == 7
    assert product["reorder_level"] == 6
    assert (movement["quantity_before"], movement["quantity_after"], movement["reason"]) == (4, 7, "Supplier receipt")


def test_dashboard_date_filter_uses_real_sales_and_links(client):
    from routes import admin_dashboard
    day = datetime.now(timezone.utc).replace(hour=12, minute=0, second=0, microsecond=0)
    admin_dashboard.sales_collection.insert_one({
        "sale_date": day, "total_revenue": 25, "profit_per_sale": 7,
        "product_name": "Real item", "quantity_sold": 2,
    })
    signed_in(client)
    response = client.get(f"/admin?start={day.date().isoformat()}&end={day.date().isoformat()}")
    assert response.status_code == 200
    assert b"Rs. 25.00" in response.data
    assert b"Real item" in response.data


def test_forecast_series_separates_history_from_four_week_estimate():
    from forecasting import forecast_sales
    start = datetime(2025, 1, 6, tzinfo=timezone.utc)
    records = [{"sale_date": start + timedelta(weeks=index), "total_revenue": 100, "profit_per_sale": 20}
               for index in range(10)]
    forecast = forecast_sales(records)
    assert forecast["available"]
    assert len(forecast["history"]) == 10
    assert len(forecast["forecast_series"]) == 4
    assert forecast["forecast_series"][0]["date"] > forecast["history"][-1]["date"]


def test_daily_sales_summary_export_groups_real_dated_records(client):
    from routes import reports
    first = datetime(2025, 4, 2, 9, tzinfo=timezone.utc)
    reports.sales_collection.insert_many([
        {"sale_date": first, "total_revenue": 20, "discount_applied": 1, "profit_per_sale": 7},
        {"sale_date": first.replace(hour=17), "total_revenue": 30, "discount_applied": 0, "profit_per_sale": 10},
        {"sale_date": first+timedelta(days=1), "total_revenue": 50, "discount_applied": 0, "profit_per_sale": 20},
    ])
    signed_in(client)
    response = client.get("/reports/export?type=daily&format=excel&start=2025-04-02&end=2025-04-02")
    assert response.status_code == 200
    rows = list(load_workbook(io.BytesIO(response.data), read_only=True).active.values)
    assert rows[1] == ("2025-04-02", 2, 50, 1, 17)


def test_every_post_form_has_a_csrf_token_and_signup_hides_admin_role():
    root = Path(__file__).resolve().parents[1]
    for path in (root / "templates").glob("*.html"):
        text = path.read_text(encoding="utf-8")
        for form in re.findall(r"<form\b[^>]*>.*?</form>", text, flags=re.I | re.S):
            if re.search(r"\bmethod\s*=\s*['\"]?post\b", form, flags=re.I):
                assert 'name="csrf_token"' in form, path.name
    register = (root / "templates" / "register.html").read_text(encoding="utf-8")
    assert 'value="admin"' not in register
